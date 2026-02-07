# core/result_exporter.py
# -*- coding: utf-8 -*-
"""
結果匯出模組

將分析結果匯出為 Excel (.xlsx) 或 CSV 格式。
每次發球事件一行，包含：
- 基本資訊（從檔名解析）
- 發球分析（現有）
- 發球區域（新增）
- 接球分析（新增）
- 品質指標
"""

import csv
import json
import os
import logging
from typing import Dict, List, Any, Optional


# 品質等級定義
def compute_quality_grade(ball_detection_rate: float) -> str:
    """
    計算品質等級

    Args:
        ball_detection_rate: 球偵測率 (0-1)

    Returns:
        'A', 'B', 'C', 或 'F'
    """
    if ball_detection_rate > 0.7:
        return 'A'
    elif ball_detection_rate > 0.5:
        return 'B'
    elif ball_detection_rate > 0.3:
        return 'C'
    else:
        return 'F'


# CSV/Excel 欄位定義
COLUMNS = [
    # 基本資訊
    'video_name',
    'venue',
    'year',
    'court',
    'gender',
    'round',
    'match_number',
    'star_level',
    'group_key',

    # 發球分析
    'serve_detected',
    'toss_frame',
    'hit_frame',
    'hit_speed',
    'server_index',
    'confidence',
    'serve_type',
    'is_jump_serve',
    'jump_height',

    # 發球區域
    'serve_zone',
    'serving_side',

    # 接球分析
    'reception_detected',
    'reception_frame',
    'reception_zone',
    'receiver_index',
    'time_to_reception',
    'reception_confidence',
    'ball_crossed_net',

    # 品質指標
    'quality_grade',
    'ball_detection_rate',
    'status',
]


def build_result_row(
    video_name: str,
    parsed_filename: Optional[Dict],
    serve_result: Optional[Dict],
    reception_result: Optional[Dict],
    ball_detection_rate: float = 0.0,
    status: str = 'success',
) -> Dict[str, Any]:
    """
    建立一行結果資料

    Args:
        video_name: 影片檔名
        parsed_filename: filename_parser 的解析結果
        serve_result: 發球分析結果
        reception_result: 接球分析結果
        ball_detection_rate: 球偵測率
        status: 處理狀態 ('success', 'no_serve', 'error', 'skipped')

    Returns:
        結果字典
    """
    row = {col: None for col in COLUMNS}

    # 基本資訊
    row['video_name'] = video_name
    row['status'] = status
    row['ball_detection_rate'] = round(ball_detection_rate, 3) if ball_detection_rate else 0
    row['quality_grade'] = compute_quality_grade(ball_detection_rate)

    if parsed_filename:
        row['venue'] = parsed_filename.get('venue')
        row['year'] = parsed_filename.get('year')
        row['court'] = parsed_filename.get('court')
        row['gender'] = parsed_filename.get('gender')
        row['round'] = parsed_filename.get('round')
        row['match_number'] = parsed_filename.get('match_number')
        row['star_level'] = parsed_filename.get('star_level')
        row['group_key'] = parsed_filename.get('group_key')

    if serve_result:
        row['serve_detected'] = serve_result.get('serve_detected', False)
        row['toss_frame'] = serve_result.get('toss_frame')
        row['hit_frame'] = serve_result.get('hit_frame')
        row['hit_speed'] = serve_result.get('hit_speed')
        row['server_index'] = serve_result.get('server_index')
        row['confidence'] = serve_result.get('confidence')
        row['serve_type'] = serve_result.get('serve_type')
        row['is_jump_serve'] = serve_result.get('is_jump_serve')
        row['jump_height'] = serve_result.get('jump_height')
        row['serve_zone'] = serve_result.get('serve_zone')
        row['serving_side'] = serve_result.get('serving_side')

    if reception_result:
        row['reception_detected'] = reception_result.get('reception_detected', False)
        row['reception_frame'] = reception_result.get('reception_frame')
        row['reception_zone'] = reception_result.get('reception_zone')
        row['receiver_index'] = reception_result.get('receiver_index')
        row['time_to_reception'] = reception_result.get('time_to_reception')
        row['reception_confidence'] = reception_result.get('confidence')
        row['ball_crossed_net'] = reception_result.get('ball_crossed_net')

    return row


def export_to_csv(results: List[Dict], output_path: str) -> str:
    """
    匯出結果到 CSV

    Args:
        results: build_result_row() 產生的結果列表
        output_path: 輸出路徑

    Returns:
        實際輸出路徑
    """
    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)

    with open(output_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction='ignore')
        writer.writeheader()
        for row in results:
            writer.writerow(row)

    logging.info(f"Exported {len(results)} results to {output_path}")
    return output_path


def export_to_excel(results: List[Dict], output_path: str) -> str:
    """
    匯出結果到 Excel (.xlsx)

    需要 openpyxl 套件。如果不可用，自動降級為 CSV。

    Args:
        results: build_result_row() 產生的結果列表
        output_path: 輸出路徑

    Returns:
        實際輸出路徑
    """
    try:
        import openpyxl
    except ImportError:
        logging.warning("openpyxl not installed, falling back to CSV export")
        csv_path = output_path.replace('.xlsx', '.csv')
        return export_to_csv(results, csv_path)

    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Serve Analysis"

    # 寫入標題
    ws.append(COLUMNS)

    # 寫入資料
    for row in results:
        ws.append([row.get(col) for col in COLUMNS])

    # 自動調整欄寬
    for col_idx, col_name in enumerate(COLUMNS, 1):
        max_len = len(col_name)
        for row_idx in range(2, min(len(results) + 2, 50)):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 2, 30)

    wb.save(output_path)
    logging.info(f"Exported {len(results)} results to {output_path}")
    return output_path


def export_summary_json(results: List[Dict], output_path: str) -> str:
    """
    匯出摘要 JSON（統計資訊）

    Args:
        results: 結果列表
        output_path: 輸出路徑

    Returns:
        實際輸出路徑
    """
    total = len(results)
    if total == 0:
        summary = {"total": 0, "message": "No results"}
    else:
        serve_detected = sum(1 for r in results if r.get('serve_detected'))
        reception_detected = sum(1 for r in results if r.get('reception_detected'))
        jump_serves = sum(1 for r in results if r.get('is_jump_serve'))

        grades = {}
        for r in results:
            g = r.get('quality_grade', 'F')
            grades[g] = grades.get(g, 0) + 1

        # 發球區分布
        serve_zone_dist = {}
        for r in results:
            z = r.get('serve_zone')
            if z is not None:
                serve_zone_dist[str(z)] = serve_zone_dist.get(str(z), 0) + 1

        # 接球區分布
        reception_zone_dist = {}
        for r in results:
            z = r.get('reception_zone')
            if z is not None:
                reception_zone_dist[str(z)] = reception_zone_dist.get(str(z), 0) + 1

        summary = {
            "total_videos": total,
            "serve_detected": serve_detected,
            "serve_detection_rate": round(serve_detected / total, 3),
            "reception_detected": reception_detected,
            "reception_detection_rate": round(reception_detected / total, 3) if total > 0 else 0,
            "jump_serves": jump_serves,
            "standing_serves": serve_detected - jump_serves,
            "quality_grades": grades,
            "serve_zone_distribution": serve_zone_dist,
            "reception_zone_distribution": reception_zone_dist,
        }

    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)

    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(summary, f, indent=2, ensure_ascii=False)

    return output_path
